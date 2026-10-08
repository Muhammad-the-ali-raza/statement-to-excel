"""Statement to Excel - turns a HABIBMETRO SMS-alert .txt file into a formatted Excel workbook."""
import datetime
import os
import re
import tkinter as tk
from tkinter import filedialog, messagebox

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

LINE_RE = re.compile(
    r"Alert:\s*(\d{2}-\d{2}-\d{2,4})\s+(\d{1,2}:\d{2}(?::\d{2})?)\s+(\S+)\s+"
    r"Amt\s+([\d,]+(?:\.\d+)?)\s+(Debited|Credited)\s*/?\s*(.*)$",
    re.IGNORECASE,
)

FONT = "Arial"
MONEY = '#,##0.00;(#,##0.00);"-"'
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", start_color="1F4E78")
TOTAL_FILL = PatternFill("solid", start_color="DDEBF7")
CREDIT_FILL = PatternFill("solid", start_color="E2EFDA")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def read_text(path):
    raw = open(path, "rb").read()
    for enc in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            text = raw.decode(enc)
            if "Alert" in text:
                return text
        except UnicodeDecodeError:
            pass
    return raw.decode("utf-8", errors="replace")


def categorize(desc, typ):
    d = desc.lower()
    if "gsm" in d:
        return "GSM Subscription Charges"
    if "sales tax" in d:
        return "Sindh Sales Tax (15%)"
    if "visa premium" in d:
        return "Visa Premium Annual Charges"
    if typ == "Credited":
        return "Incoming Transfer"
    return desc[:40] or "Other"


def parse(text):
    rows, skipped = [], []
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        m = LINE_RE.search(line)
        if not m:
            skipped.append(line)
            continue
        d, t, acct, amt, typ, desc = m.groups()
        typ = typ.capitalize()
        desc = re.sub(r"\s*\.\.\.$", "", desc)
        desc = re.sub(r"\s*\S?Monthly\s*\S?$", " [Monthly]", desc).strip()
        date = datetime.datetime.strptime(d, "%d-%m-%Y" if len(d) == 10 else "%d-%m-%y").date()
        time = datetime.datetime.strptime(t, "%H:%M:%S" if t.count(":") == 2 else "%H:%M").time()
        rows.append((date, time, acct, desc, categorize(desc, typ), typ, float(amt.replace(",", ""))))
    return rows, skipped


def header_row(ws, row, heads):
    for c, h in enumerate(heads, 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def build_workbook(rows, src_name, out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws["A1"] = "HABIBMETRO Account Statement"
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    ws["A2"] = f"Source: SMS alerts in '{src_name}' (descriptions may be truncated in the original alerts)"
    ws["A2"].font = Font(name=FONT, italic=True, color="808080")
    header_row(ws, 4, ["S.No", "Date", "Time", "Account", "Description", "Category", "Type",
                       "Debit (PKR)", "Credit (PKR)", "Running Balance (PKR)"])

    start = 5
    for i, (date, time, acct, desc, cat, typ, amt) in enumerate(rows):
        r = start + i
        vals = [i + 1, date, time, acct, desc, cat, typ,
                amt if typ == "Debited" else None,
                amt if typ == "Credited" else None,
                f"=I{r}-H{r}" if i == 0 else f"=J{r-1}+I{r}-H{r}"]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = Font(name=FONT)
            cell.border = BORDER
            if typ == "Credited":
                cell.fill = CREDIT_FILL
        ws.cell(row=r, column=2).number_format = "DD-MMM-YYYY"
        ws.cell(row=r, column=3).number_format = "HH:MM"
        for c in (8, 9, 10):
            ws.cell(row=r, column=c).number_format = MONEY
    end = start + len(rows) - 1

    tr = end + 1
    ws.cell(row=tr, column=7, value="TOTAL")
    ws.cell(row=tr, column=8, value=f"=SUM(H{start}:H{end})")
    ws.cell(row=tr, column=9, value=f"=SUM(I{start}:I{end})")
    ws.cell(row=tr, column=10, value=f"=I{tr}-H{tr}")
    for c in range(1, 11):
        cell = ws.cell(row=tr, column=c)
        cell.font = Font(name=FONT, bold=True)
        cell.border = BORDER
        cell.fill = TOTAL_FILL
        if c >= 8:
            cell.number_format = MONEY
    for col, w in zip("ABCDEFGHIJ", [7, 13, 8, 13, 48, 28, 10, 14, 14, 20]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:J{end}"

    s = wb.create_sheet("Summary")
    s["A1"] = "Summary by Category"
    s["A1"].font = Font(name=FONT, bold=True, size=14)
    header_row(s, 3, ["Category", "No. of Transactions", "Total Debit (PKR)", "Total Credit (PKR)"])
    cats = list(dict.fromkeys(r[4] for r in sorted(rows, key=lambda r: r[5] != "Credited")))
    rng = lambda col: f"Transactions!${col}${start}:${col}${end}"
    for i, cat in enumerate(cats):
        r = 4 + i
        s.cell(row=r, column=1, value=cat)
        s.cell(row=r, column=2, value=f"=COUNTIF({rng('F')},A{r})")
        s.cell(row=r, column=3, value=f"=SUMIF({rng('F')},A{r},{rng('H')})")
        s.cell(row=r, column=4, value=f"=SUMIF({rng('F')},A{r},{rng('I')})")
    last = 4 + len(cats) - 1
    t = last + 1
    s.cell(row=t, column=1, value="TOTAL")
    for c, col in ((2, "B"), (3, "C"), (4, "D")):
        s.cell(row=t, column=c, value=f"=SUM({col}4:{col}{last})")
    for r in range(4, t + 1):
        for c in range(1, 5):
            cell = s.cell(row=r, column=c)
            cell.border = BORDER
            cell.font = Font(name=FONT, bold=(r == t))
            if r == t:
                cell.fill = TOTAL_FILL
            if c >= 3:
                cell.number_format = MONEY
    s.cell(row=t + 2, column=1, value="Net (Credit - Debit)").font = Font(name=FONT, bold=True)
    net = s.cell(row=t + 2, column=4, value=f"=D{t}-C{t}")
    net.font = Font(name=FONT, bold=True)
    net.number_format = MONEY
    for col, w in zip("ABCD", [42, 20, 20, 20]):
        s.column_dimensions[col].width = w

    wb.save(out_path)


def convert():
    src = filedialog.askopenfilename(title="Select statement text file",
                                     filetypes=[("Text files", "*.txt"), ("All files", "*.*")])
    if not src:
        return
    rows, skipped = parse(read_text(src))
    if not rows:
        messagebox.showerror("No transactions found",
                             "Could not find any HABIBMETRO alert lines in this file.")
        return

    out = os.path.splitext(src)[0] + ".xlsx"
    try:
        build_workbook(rows, os.path.basename(src), out)
    except PermissionError:
        messagebox.showerror("File is open",
                             f"Please close this file in Excel and try again:\n\n{out}")
        return

    debit = sum(r[6] for r in rows if r[5] == "Debited")
    credit = sum(r[6] for r in rows if r[5] == "Credited")
    msg = (f"Excel file created:\n{out}\n\n"
           f"Transactions: {len(rows)}\n"
           f"Total credit: PKR {credit:,.2f}\n"
           f"Total debit:  PKR {debit:,.2f}\n"
           f"Net:          PKR {credit - debit:,.2f}")
    if skipped:
        msg += f"\n\nNote: {len(skipped)} line(s) could not be read and were skipped."
    status.config(text=f"Done: {os.path.basename(out)}")
    if messagebox.askyesno("Done", msg + "\n\nOpen the Excel file now?"):
        os.startfile(out)


root = tk.Tk()
root.title("Statement to Excel")
root.geometry("420x200")
root.resizable(False, False)
tk.Label(root, text="Bank SMS Statement to Excel", font=("Arial", 14, "bold")).pack(pady=(25, 5))
tk.Label(root, text="Choose your .txt file. The Excel file is saved next to it.",
         font=("Arial", 9)).pack()
tk.Button(root, text="Select TXT File", font=("Arial", 11, "bold"), bg="#1F4E78", fg="white",
          padx=20, pady=6, command=convert).pack(pady=18)
status = tk.Label(root, text="", font=("Arial", 9), fg="#2E7D32")
status.pack()
root.mainloop()
